"""Deep analysis of template PPTX layouts, placeholders, fonts, colors, and all design elements."""
import sys, pathlib
sys.path.insert(0, str(pathlib.Path(__file__).parent))

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.oxml.ns import qn

TRAIN_DIR = pathlib.Path(__file__).parent / "train"

for tmpl_path in sorted(TRAIN_DIR.glob("*.pptx")):
    prs = Presentation(str(tmpl_path))
    print(f"\n{'='*80}")
    print(f"  TEMPLATE: {tmpl_path.name}")
    print(f"  Slide size: {prs.slide_width/914400:.2f}\" x {prs.slide_height/914400:.2f}\"")
    print(f"  Layouts: {len(prs.slide_layouts)}")
    print(f"  Existing slides: {len(prs.slides)}")
    print(f"{'='*80}")

    # Extract theme colors
    print("\n  THEME COLORS:")
    try:
        theme = prs.slide_masters[0].element.find(qn("p:cSld"))
        theme_elem = prs.slide_masters[0].element
        clrScheme = theme_elem.find('.//' + qn('a:clrScheme'))
        if clrScheme is not None:
            for child in clrScheme:
                tag = child.tag.split('}')[-1] if '}' in child.tag else child.tag
                srgb = child.find(qn('a:srgbClr'))
                sys_clr = child.find(qn('a:sysClr'))
                if srgb is not None:
                    print(f"    {tag:15} = #{srgb.get('val')}")
                elif sys_clr is not None:
                    print(f"    {tag:15} = system:{sys_clr.get('val')} (last=#{sys_clr.get('lastClr', '?')})")
        else:
            print("    No color scheme found")
    except Exception as e:
        print(f"    Error: {e}")

    # Extract theme fonts
    print("\n  THEME FONTS:")
    try:
        fontScheme = prs.slide_masters[0].element.find('.//' + qn('a:fontScheme'))
        if fontScheme is not None:
            print(f"    Scheme name: {fontScheme.get('name', '?')}")
            for ftype in ['majorFont', 'minorFont']:
                font_elem = fontScheme.find(qn(f'a:{ftype}'))
                if font_elem is not None:
                    latin = font_elem.find(qn('a:latin'))
                    if latin is not None:
                        print(f"    {ftype:12} = {latin.get('typeface', '?')}")
    except Exception as e:
        print(f"    Error: {e}")

    # Analyze each layout
    print(f"\n  LAYOUTS:")
    for i, layout in enumerate(prs.slide_layouts):
        name = layout.name if hasattr(layout, 'name') else f"Layout {i}"
        shapes = list(layout.shapes)
        ph_info = []
        for shape in shapes:
            if shape.is_placeholder:
                t = shape.placeholder_format.type
                idx = shape.placeholder_format.idx
                left = shape.left / 914400 if shape.left else 0
                top = shape.top / 914400 if shape.top else 0
                w = shape.width / 914400 if shape.width else 0
                h = shape.height / 914400 if shape.height else 0
                ph_info.append(f"      ph[{idx}] type={t} pos=({left:.1f},{top:.1f}) size=({w:.1f}x{h:.1f})")
        print(f"    [{i}] \"{name}\" — {len(shapes)} shapes, {len(ph_info)} placeholders")
        for info in ph_info:
            print(info)

    # Analyze existing slides to see what the template actually looks like
    print(f"\n  EXISTING SLIDES:")
    for si, slide in enumerate(prs.slides):
        layout_name = slide.slide_layout.name if hasattr(slide.slide_layout, 'name') else '?'
        shapes = list(slide.shapes)
        text_preview = ""
        for shape in shapes:
            if shape.has_text_frame:
                t = shape.text_frame.text[:60]
                if t.strip():
                    text_preview = t
                    break
        print(f"    Slide {si}: layout=\"{layout_name}\" shapes={len(shapes)} text=\"{text_preview}\"")
    print()
