"""Advanced Template Analyzer to extract colors from shapes instead of just theme XML."""
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
import pathlib
import json

def extract_colors_from_shape(shape):
    colors = []
    # Fill colors
    try:
        if shape.fill.type == 1: # Solid fill
            rgb = shape.fill.fore_color.rgb
            if rgb: colors.append(f"#{rgb}")
    except: pass
    
    # Line colors
    try:
        rgb = shape.line.color.rgb
        if rgb: colors.append(f"#{rgb}")
    except: pass
    
    # Text colors
    if shape.has_text_frame:
        for p in shape.text_frame.paragraphs:
            for r in p.runs:
                try:
                    rgb = r.font.color.rgb
                    if rgb: colors.append(f"#{rgb}")
                except: pass
    return colors

def analyze(tmpl_file):
    prs = Presentation(tmpl_file)
    extracted_colors = set()
    
    # Master slides
    for master in prs.slide_masters:
        if master.background and hasattr(master.background, 'fill') and master.background.fill.type == 1:
             try:
                 extracted_colors.add(f"#{master.background.fill.fore_color.rgb}")
             except: pass
        for shape in master.shapes:
            extracted_colors.update(extract_colors_from_shape(shape))
            
    # Layouts
    for layout in prs.slide_layouts:
        for shape in layout.shapes:
            extracted_colors.update(extract_colors_from_shape(shape))
            
    print(f"\nExtracted from {tmpl_file.name}:")
    print(list(extracted_colors))

TRAIN_DIR = pathlib.Path(__file__).parent / "train"
for f in TRAIN_DIR.glob("*.pptx"):
    analyze(f)
