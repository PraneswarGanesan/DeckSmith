from pptx import Presentation


def load_template(path: str):
    return Presentation(path)


def replace_text(slide, placeholder, value):
    for shape in slide.shapes:
        if shape.has_text_frame:
            if placeholder in shape.text:
                shape.text = shape.text.replace(placeholder, value)


def insert_image(slide, image_bytes, left, top, width):
    import io
    image_stream = io.BytesIO(image_bytes)
    slide.shapes.add_picture(image_stream, left, top, width=width)