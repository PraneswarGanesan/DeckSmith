from pptx import Presentation
from pptx.util import Inches
import io


def build_ppt(template_path, slides, charts=None):
    prs = Presentation(template_path)

    for i, s in enumerate(slides):

        slide = prs.slides.add_slide(prs.slide_layouts[1])

        if slide.shapes.title:
            slide.shapes.title.text = s["title"]

        for shape in slide.placeholders:
            if shape.placeholder_format.type in [2, 7]:
                tf = shape.text_frame
                tf.clear()

                for bullet in s["content"]:
                    p = tf.add_paragraph()
                    p.text = bullet

        if charts and s.get("chart"):
            img = charts.pop(0)
            stream = io.BytesIO(img)
            slide.shapes.add_picture(stream, Inches(5), Inches(2), width=Inches(4))

    return prs