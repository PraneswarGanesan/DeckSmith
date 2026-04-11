"""
Template Analyzer Agent — Scans the uploaded PPTX template and extracts semantics.

Design Philosophy:
  Before styling a presentation, we must know WHAT shapes/layouts the user's
  template provides. We extract this schema and pass it down the pipeline.
  This allows the Visual Composer to pick native layouts rather than drawing
  primitive rectangles that clash with the template's aesthetics.
"""
from __future__ import annotations

import io
from pptx import Presentation
from config import settings
from logger import get_logger

logger = get_logger(__name__)

# Map pptx placeholder enum strings to simpler concepts
_PH_MAP = {
    "TITLE": "title",
    "CENTER_TITLE": "title",
    "SUBTITLE": "subtitle",
    "BODY": "body",
    "OBJECT": "object",    # often used for charts/tables/text
    "PICTURE": "image",
    "TABLE": "table",
    "CHART": "chart",
}

def _simplify_type(ph_type_str: str) -> str:
    """Convert 'PP_PLACEHOLDER.OBJECT (7)' -> 'object'"""
    base = ph_type_str.split(' ')[0].split('.')[-1]
    return _PH_MAP.get(base, "unknown")


def _analyze_layouts(prs: Presentation) -> list[dict]:
    """
    Extract a clean schema of all SlideLayouts inside a template.
    Returns something like:
    [
      { "layout_index": 0, "name": "Title Slide", "placeholders": [{"idx": 0, "type": "title"}] },
      { "layout_index": 1, "name": "3 Column", "placeholders": [{"idx": 0, "type": "title"}, {"idx":1, "type":"body"}, {"idx":2, "type":"body"}, {"idx":3, "type":"body"}] }
    ]
    """
    layouts = []
    
    for i, layout in enumerate(prs.slide_layouts):
        l_info = {
            "layout_index": i,
            "name": layout.name,
            "placeholders": []
        }
        
        # Only care about placeholders, ignore hardcoded backgrounds or lines
        for shape in layout.shapes:
            if shape.is_placeholder:
                ph = shape.placeholder_format
                ph_type = _simplify_type(str(ph.type))
                l_info["placeholders"].append({
                    "idx": ph.idx,
                    "type": ph_type
                })
                
        # We only consider layouts that have at least one placeholder
        if len(l_info["placeholders"]) > 0:
            layouts.append(l_info)
            
    return layouts


async def template_analyzer_node(state: dict) -> dict:
    """
    Downloads the blank template, analyzes its layout schema, and pushes it
    into the state so the Visual Composer knows what's possible.
    """
    template_id = state.get("template_id")
    if not template_id:
        logger.warning("[TemplateAnalyzer] No template_id passed. Providing empty schema.")
        return {"template_schema": []}
        
    logger.info(f"[TemplateAnalyzer] Analyzing template layout schemas for {template_id}...")
    
    try:
        from core.database import get_template
        from services.storage_service import download_file
        
        tmpl_meta = get_template(template_id)
        tmpl_bytes = download_file(settings.STORAGE_BUCKET_RAW, tmpl_meta["storage_path"])
        prs = Presentation(io.BytesIO(tmpl_bytes))
        
        schema = _analyze_layouts(prs)
        logger.info(f"[TemplateAnalyzer] Found {len(schema)} viable slide layouts in template.")
        return {"template_schema": schema}
        
    except Exception as exc:
        logger.error(f"[TemplateAnalyzer] Failed to analyze template {template_id}: {exc}")
        return {"template_schema": []}
